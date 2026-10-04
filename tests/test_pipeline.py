"""End to end on the sample data: every number in the report must reconcile with the order lines."""

import pytest
from openpyxl import load_workbook

from salesreport.cli import main
from salesreport.pipeline import run


@pytest.fixture
def result(config_path, sample_dir, tmp_path, fetch):
    output = tmp_path / "report.xlsx"
    res = run(config_path, sample_dir, output, tmp_path / "rates.json", fetch=fetch)
    return res, load_workbook(output)


def column(ws, header):
    names = [c.value for c in ws[1]]
    return [row[names.index(header)] for row in ws.iter_rows(min_row=2, values_only=True)]


def test_every_row_is_accounted_for(result):
    res, _ = result
    q = res.quality
    assert q.rows_read == len(res.lines) + q.duplicates + q.rejected
    assert q.rejected == 6  # the rows broken on purpose in make_sample_data.py
    assert q.duplicates > 0


def test_summary_matches_order_lines(result):
    res, wb = result
    orders_total = round(sum(column(wb["Orders"], "Amount EUR")), 2)
    summary = wb["Summary"]
    assert summary["A4"].value == "Net revenue (EUR)"
    assert summary["A5"].value == pytest.approx(orders_total, abs=0.01)
    monthly_total = next(r for r in summary.iter_rows(values_only=True) if r[0] == "Total")
    assert monthly_total[4] == pytest.approx(orders_total, abs=0.01)
    assert res.kpis["Orders"] == len(
        {(c, o) for c, o in zip(res.lines.channel, res.lines.order_id)}
    )


def test_foreign_currency_is_converted(result):
    _, wb = result
    ws = wb["Orders"]
    for currency, amount, rate, eur in zip(
        column(ws, "Currency"),
        column(ws, "Amount"),
        column(ws, "FX rate (EUR=1)"),
        column(ws, "Amount EUR"),
    ):
        expected = {"EUR": 1.0, "GBP": 0.85, "USD": 1.10}[currency]
        assert rate == expected
        assert eur == pytest.approx(amount / rate, abs=0.005)


def test_issues_sheet_lists_rejections_and_duplicates(result):
    res, wb = result
    problems = column(wb["Issues"], "Problem")
    assert len(problems) == res.quality.rejected + res.quality.duplicates
    assert "invalid date '31/09/2026'" in problems
    assert sum(p.startswith("duplicate of") for p in problems) == res.quality.duplicates


def test_cli_reports_a_missing_folder_with_exit_code_1(tmp_path, config_path, caplog):
    code = main(["--config", str(config_path), "--input", str(tmp_path / "nope"), "--offline"])
    assert code == 1
    assert "input folder not found" in caplog.text
