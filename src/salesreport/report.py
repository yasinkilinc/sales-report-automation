"""Write the formatted Excel report: Summary (KPIs, monthly table, chart, top products), Orders, Issues."""

from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ACCENT = "1F6F5C"
HEADER_FILL = PatternFill("solid", fgColor=ACCENT)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_BORDER = Border(top=Side(style="thin", color="000000"))
MONEY = "#,##0.00"
COUNT = "#,##0"


@dataclass(frozen=True)
class Quality:
    rows_read: int
    duplicates: int
    rejected: int


def write_report(
    lines: pd.DataFrame, issues: pd.DataFrame, base: str, quality: Quality, path: Path
):
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    _summary_sheet(summary, lines, base, quality)
    _orders_sheet(wb.create_sheet("Orders"), lines, base)
    _issues_sheet(wb.create_sheet("Issues"), issues)
    for ws in wb.worksheets:  # print: landscape, all columns on one page width
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def kpis(lines: pd.DataFrame, base: str) -> dict[str, float]:
    amount = lines[f"amount_{base.lower()}"]
    orders = lines[["channel", "order_id"]].drop_duplicates().shape[0]
    revenue = round(amount.sum(), 2)
    return {
        f"Net revenue ({base})": revenue,
        "Orders": orders,
        "Units sold": int(lines.loc[lines["quantity"] > 0, "quantity"].sum()),
        f"Average order ({base})": round(revenue / orders, 2) if orders else 0.0,
        f"Returns ({base})": round(amount[lines["quantity"] < 0].sum(), 2),
    }


def monthly(lines: pd.DataFrame, base: str) -> pd.DataFrame:
    """Revenue per month (rows) and channel (columns), with a Total column."""
    month = lines["date"].dt.to_period("M").rename("Month")
    table = lines.pivot_table(
        index=month, columns="channel", values=f"amount_{base.lower()}", aggfunc="sum", fill_value=0
    ).round(2)
    table["Total"] = table.sum(axis=1).round(2)
    table.index = table.index.strftime("%b %Y")
    return table


def top_products(lines: pd.DataFrame, base: str, n: int = 10) -> pd.DataFrame:
    grouped = lines.groupby(["sku", "product"], as_index=False).agg(
        units=("quantity", "sum"), revenue=(f"amount_{base.lower()}", "sum")
    )
    grouped["revenue"] = grouped["revenue"].round(2)
    return grouped.sort_values("revenue", ascending=False).head(n).reset_index(drop=True)


def _summary_sheet(ws: Worksheet, lines: pd.DataFrame, base: str, quality: Quality):
    start, end = lines["date"].min(), lines["date"].max()
    ws["A1"] = "Sales report"
    ws["A1"].font = Font(bold=True, size=18, color=ACCENT)
    ws["A2"] = f"{start:%d %b %Y} to {end:%d %b %Y}. Amounts in {base} at ECB reference rates."
    ws["A2"].font = Font(italic=True, color="555555")

    for col, (label, value) in enumerate(kpis(lines, base).items(), start=1):
        ws.cell(4, col, label).font = Font(color="555555")
        cell = ws.cell(5, col, value)
        cell.font = Font(bold=True, size=14)
        cell.number_format = COUNT if isinstance(value, int) else MONEY

    month_table = monthly(lines, base)
    last_row = _write_table(ws, 7, f"Revenue by month ({base})", month_table, total_row=True)
    _add_chart(
        ws, month_table, header_row=8, anchor=f"{get_column_letter(len(month_table.columns) + 3)}7"
    )

    products = top_products(lines, base).rename(
        columns={
            "sku": "SKU",
            "product": "Product",
            "units": "Units",
            "revenue": f"Revenue ({base})",
        }
    )
    last_row = _write_table(ws, last_row + 2, "Top 10 products", products.set_index("SKU"))

    ws.cell(last_row + 2, 1, "Data quality").font = Font(bold=True, size=12)
    checks = [("Rows read", quality.rows_read), ("Duplicates removed", quality.duplicates)]
    checks.append(("Rows rejected (see Issues sheet)", quality.rejected))
    for offset, (label, value) in enumerate(checks, start=3):
        ws.cell(last_row + offset, 1, label)
        ws.cell(last_row + offset, 2, value).number_format = COUNT
    _fit_columns(ws, min_width=14)
    ws.column_dimensions["A"].width = 30


def _write_table(ws: Worksheet, row: int, title: str, table: pd.DataFrame, total_row=False) -> int:
    """Writes title, header and rows starting at `row`. Returns the last row used."""
    ws.cell(row, 1, title).font = Font(bold=True, size=12)
    header = [table.index.name or "Month", *table.columns]
    for col, name in enumerate(header, start=1):
        _header(ws.cell(row + 1, col, name))
    r = row + 1
    for r, (key, values) in enumerate(table.iterrows(), start=row + 2):
        ws.cell(r, 1, key)
        for col, value in enumerate(values, start=2):
            _value(ws.cell(r, col, value))
    if total_row:
        r += 1
        ws.cell(r, 1, "Total").font = Font(bold=True)
        for col, value in enumerate(table.sum().round(2), start=2):
            cell = ws.cell(r, col, value)
            _value(cell)
            cell.font = Font(bold=True)
        for col in range(1, len(header) + 1):
            ws.cell(r, col).border = TOTAL_BORDER
    return r


def _add_chart(ws: Worksheet, table: pd.DataFrame, header_row: int, anchor: str):
    channels = len(table.columns) - 1  # without Total
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "Revenue by month and channel"
    chart.height, chart.width = 8, 16
    chart.y_axis.numFmt = "#,##0"
    chart.x_axis.delete = chart.y_axis.delete = False  # openpyxl hides axes in Excel otherwise
    chart.legend.position = "b"
    data = Reference(
        ws, min_col=2, max_col=1 + channels, min_row=header_row, max_row=header_row + len(table)
    )
    months = Reference(ws, min_col=1, min_row=header_row + 1, max_row=header_row + len(table))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(months)
    ws.add_chart(chart, anchor)


def _orders_sheet(ws: Worksheet, lines: pd.DataFrame, base: str):
    columns = {
        "date": "Date", "channel": "Channel", "order_id": "Order ID", "sku": "SKU",
        "product": "Product", "quantity": "Qty", "unit_price": "Unit price", "currency": "Currency",
        "amount": "Amount", "fx_rate": f"FX rate ({base}=1)", f"amount_{base.lower()}": f"Amount {base}",
        "source_file": "Source file", "source_row": "Source row",
    }  # fmt: skip
    rows = lines.sort_values(["date", "channel", "order_id"])[list(columns)]
    _write_sheet(ws, rows.rename(columns=columns))
    for col, fmt in {"A": "yyyy-mm-dd", "G": MONEY, "I": MONEY, "J": "0.0000", "K": MONEY}.items():
        for cell in ws[col][1:]:
            cell.number_format = fmt


def _issues_sheet(ws: Worksheet, issues: pd.DataFrame):
    columns = {
        "channel": "Channel", "source_file": "Source file", "source_row": "Row",
        "problem": "Problem", "order_id": "Order ID", "sku": "SKU",
    }  # fmt: skip
    # Rejected rows first: they need a decision, duplicates are only informational
    ordered = issues.assign(_dup=issues["problem"].str.startswith("duplicate of"))
    ordered = ordered.sort_values(["_dup", "source_file", "source_row"]).drop(columns="_dup")
    _write_sheet(ws, ordered.rename(columns=columns))
    for cell in ws["C"][1:]:
        cell.alignment = Alignment(horizontal="center")


def _write_sheet(ws: Worksheet, frame: pd.DataFrame):
    ws.append(list(frame.columns))
    for values in frame.itertuples(index=False):
        ws.append([None if pd.isna(v) else v for v in values])
    for cell in ws[1]:
        _header(cell)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _fit_columns(ws)


def _header(cell):
    cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _value(cell):
    if hasattr(cell.value, "item"):  # numpy scalar -> plain Python number
        cell.value = cell.value.item()
    cell.number_format = COUNT if isinstance(cell.value, int) else MONEY


def _fit_columns(ws: Worksheet, min_width: int = 8, max_width: int = 45):
    widths: dict[str, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                letter = cell.column_letter
                widths[letter] = max(widths.get(letter, 0), len(str(cell.value)))
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = min(max(width + 2, min_width), max_width)
